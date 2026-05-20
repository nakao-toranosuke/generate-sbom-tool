from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any, Optional

import openpyxl


@dataclass
class NormalizedTable:
    headers: list[str]
    rows: list[list[Any]]
    header_row_index_1based: int


def _fill_merged_cells(ws) -> None:
    # Fill merged cells (propagate top-left value)
    for merged in list(ws.merged_cells.ranges):
        min_row, min_col, max_row, max_col = merged.min_row, merged.min_col, merged.max_row, merged.max_col
        value = ws.cell(row=min_row, column=min_col).value
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                ws.cell(row=r, column=c).value = value


def normalize_excel_first_sheet(
    xlsx_path: Path,
    header_row_index_1based: Optional[int] = None,
    scan_rows: int = 30,
    header_present: bool = True,
) -> NormalizedTable:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.worksheets[0]
    _fill_merged_cells(ws)

    max_cols = ws.max_column
    max_rows = min(ws.max_row, scan_rows)

    def non_empty_count(r: int) -> int:
        cnt = 0
        for c in range(1, max_cols + 1):
            v = ws.cell(row=r, column=c).value
            if v is not None and str(v).strip() != '':
                cnt += 1
        return cnt

    if header_row_index_1based is None:
        best_r, best_cnt = 1, 0
        for r in range(1, max_rows + 1):
            cnt = non_empty_count(r)
            if cnt > best_cnt:
                best_r, best_cnt = r, cnt
        header_row_index_1based = best_r

    data_start_row = header_row_index_1based + 1
    if not header_present:
        data_start_row = header_row_index_1based

    rows: list[list[Any]] = []
    for r in range(data_start_row, ws.max_row + 1):
        row = [ws.cell(row=r, column=c).value for c in range(1, max_cols + 1)]
        if all(v is None or str(v).strip() == '' for v in row):
            continue
        rows.append(row)

    if header_present:
        headers: list[str] = []
        for c in range(1, max_cols + 1):
            v = ws.cell(row=header_row_index_1based, column=c).value
            headers.append('' if v is None else str(v).strip())
    else:
        headers = [f'col_{c}' for c in range(1, max_cols + 1)]

    def col_all_empty(idx: int) -> bool:
        if header_present and headers[idx] and str(headers[idx]).strip() != '':
            return False
        for row in rows:
            v = row[idx]
            if v is not None and str(v).strip() != '':
                return False
        return True

    trim_to = len(headers)
    while trim_to > 0 and col_all_empty(trim_to - 1):
        trim_to -= 1

    headers = headers[:trim_to]
    rows = [row[:trim_to] for row in rows]

    return NormalizedTable(headers=headers, rows=rows, header_row_index_1based=int(header_row_index_1based))
